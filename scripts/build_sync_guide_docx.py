from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "HUONG_DAN_DONG_BO_DU_LIEU.docx"
IMAGE_DIR = ROOT / "docs" / "images"

ORANGE = "F15A24"
ORANGE_DARK = "C2410C"
INK = "1F2937"
MUTED = "6B7280"
LIGHT_ORANGE = "FFF3E8"
LIGHT_GRAY = "F3F4F6"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def paragraph_rule(paragraph, color=ORANGE, size=10):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(document, marker):
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(item.get(qn("w:numId")))
        for item in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if marker == "number" else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if marker == "number" else "•")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(document, text, num_id, bold_prefix=None):
    paragraph = document.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([level, num])
    p_pr.insert(0, num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, size=11, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=INK)
    return paragraph


def add_heading(document, text, level):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 12}[level],
        bold=True,
        color=ORANGE_DARK if level < 3 else INK,
    )
    return paragraph


def add_figure(document, filename, caption, width):
    image_path = IMAGE_DIR / filename
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(width))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", caption)

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(10)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, italic=True, color=MUTED)


def add_note(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    shade_paragraph(paragraph, LIGHT_ORANGE)
    first = paragraph.add_run(f"{label}: ")
    set_run_font(first, size=10.5, bold=True, color=ORANGE_DARK)
    rest = paragraph.add_run(text)
    set_run_font(rest, size=10.5, color=INK)


def set_modern_word_compatibility(document):
    settings = document.settings.element
    for item in settings.findall(f".//{qn('w:compatSetting')}"):
        if item.get(qn("w:name")) == "compatibilityMode":
            item.set(qn("w:val"), "15")
            return


def build_document():
    document = Document()
    set_modern_word_compatibility(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, 18, 10),
        2: (13, 14, 7),
        3: (12, 10, 5),
    }
    for level, (size, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(
            ORANGE_DARK if level < 3 else INK
        )
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(2)
    header_run = header_paragraph.add_run("SUPER POST KING  |  HƯỚNG DẪN SỬ DỤNG")
    set_run_font(header_run, size=8.5, bold=True, color=MUTED)
    paragraph_rule(header_paragraph, color="E5E7EB", size=6)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run("King Automation AI  |  Trang ")
    set_run_font(footer_run, size=9, color=MUTED)
    add_page_field(footer_paragraph)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(14)
    kicker.paragraph_format.space_after = Pt(4)
    kicker_run = kicker.add_run("TÀI LIỆU HƯỚNG DẪN NHANH")
    set_run_font(kicker_run, size=9, bold=True, color=ORANGE)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("HƯỚNG DẪN ĐỒNG BỘ DỮ LIỆU")
    set_run_font(title_run, size=26, bold=True, color=INK)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run(
        "Đồng bộ giữa các nền tảng và kết nối Google Sheets"
    )
    set_run_font(subtitle_run, size=13, color=MUTED)
    paragraph_rule(subtitle, color=ORANGE, size=12)

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    meta_run = meta.add_run("Cập nhật: 17/08/2026")
    set_run_font(meta_run, size=9.5, italic=True, color=MUTED)

    add_note(
        document,
        "Có 2 cách đồng bộ",
        "Đồng bộ một cột sang nhiều nền tảng khác, hoặc đồng bộ toàn bộ dữ liệu với Google Sheets.",
    )

    number_id = add_numbering_definition(document, "number")
    bullet_id = add_numbering_definition(document, "bullet")

    add_heading(document, "1. Đồng bộ giữa các nền tảng", 1)
    intro = document.add_paragraph(
        "Dùng chức năng này khi muốn sao chép nhanh một cột như Nội dung, Media, "
        "Tên Profile hoặc Bình luận sang các nền tảng khác."
    )
    intro.paragraph_format.space_after = Pt(8)
    add_figure(
        document,
        "dong-bo-du-lieu-da-nen-tang.png",
        "Hình 1. Chọn nền tảng đích, phạm vi và cách xử lý dữ liệu.",
        5.55,
    )

    add_heading(document, "Cách thực hiện", 2)
    steps = [
        "Mở bảng dữ liệu của nền tảng nguồn, ví dụ Đăng bài Facebook.",
        "Tìm cột cần đồng bộ như Tên Profile, Nội dung, Media hoặc Bình luận.",
        "Bấm biểu tượng ↻ cạnh tên cột.",
        "Chọn một hoặc nhiều nền tảng đích.",
        "Chọn phạm vi: Dòng đã chọn, Dòng đang hiển thị hoặc Toàn bộ bảng.",
        "Chọn Ghi đè dữ liệu hiện có hoặc Chỉ điền ô trống.",
        "Bật Tạo dữ liệu còn thiếu nếu muốn tool tự tạo cột hoặc hàng cần thiết.",
        "Bấm Xem trước đồng bộ.",
    ]
    for step in steps:
        add_list_item(document, step, number_id)

    add_heading(document, "Kiểm tra trước khi ghi", 2)
    add_figure(
        document,
        "xem-truoc-dong-bo-du-lieu.png",
        "Hình 2. Tool báo số ô thay đổi và số ô được giữ nguyên.",
        5.2,
    )
    add_list_item(
        document,
        "Nếu kết quả đúng, bấm Đồng bộ ngay.",
        bullet_id,
        "Nếu kết quả đúng",
    )
    add_list_item(
        document,
        "Nếu muốn chỉnh lại, bấm Quay lại.",
        bullet_id,
        "Nếu muốn chỉnh lại",
    )
    add_list_item(
        document,
        "Bước xem trước không làm thay đổi dữ liệu.",
        bullet_id,
        "Bước xem trước",
    )
    add_note(
        document,
        "Đồng bộ ngược lại",
        "Mở nền tảng cần dùng làm nguồn rồi thực hiện lại các bước trên.",
    )

    add_heading(document, "2. Đồng bộ với Google Sheets", 1)
    add_figure(
        document,
        "dong-bo-google-sheets-cai-dat.png",
        "Hình 3. Kết nối JSON, Google Sheet và chọn hướng đồng bộ.",
        6.35,
    )

    add_heading(document, "Kết nối lần đầu", 2)
    sheets_steps = [
        "Vào Cài đặt và tìm mục Đồng bộ Google Sheets.",
        "Bấm Chọn JSON và chọn file quyền được cung cấp.",
        "Sao chép email xuất hiện trong ô Email chia sẻ Sheet.",
        "Mở Google Sheet, bấm Chia sẻ và cấp quyền Người chỉnh sửa cho email đó.",
        "Dán link Google Sheet vào ô Link hoặc ID Google Sheets.",
        "Bấm Lưu.",
        "Bấm Kiểm tra.",
        "Chỉ sử dụng khi tool báo kết nối thành công.",
    ]
    for step in sheets_steps:
        add_list_item(document, step, number_id)

    add_heading(document, "Chọn cách đồng bộ", 2)
    options = [
        ("Đẩy lên", "đưa dữ liệu từ tool lên Google Sheets."),
        ("Lấy về", "lấy dữ liệu từ Google Sheets về tool."),
        ("Hai chiều", "cập nhật dữ liệu theo cả hai hướng."),
        ("Tự động", "chọn chu kỳ và hướng để tool tự chạy."),
    ]
    for label, description in options:
        add_list_item(
            document,
            f"{label}: {description}",
            bullet_id,
            f"{label}:",
        )
    add_note(
        document,
        "Cách lưu dữ liệu",
        "Mỗi nền tảng được lưu trong một tab riêng trên Google Sheets.",
    )

    add_heading(document, "3. Lưu ý an toàn", 1)
    warnings = [
        "Nên tạo Backup trước lần đồng bộ đầu tiên.",
        "Dùng Chỉ điền ô trống nếu không muốn thay dữ liệu đang có.",
        "Luôn xem trước trước khi bấm Đồng bộ ngay.",
        "Đồng bộ tự động không tự xóa dữ liệu.",
        "Không gửi file JSON quyền cho người khác.",
    ]
    for warning in warnings:
        add_list_item(document, warning, bullet_id)

    document.core_properties.title = "Hướng dẫn đồng bộ dữ liệu - Super Post King"
    document.core_properties.subject = "Hướng dẫn sử dụng chức năng đồng bộ"
    document.core_properties.author = "King Automation AI"
    document.core_properties.keywords = "Super Post King, đồng bộ dữ liệu, Google Sheets"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
