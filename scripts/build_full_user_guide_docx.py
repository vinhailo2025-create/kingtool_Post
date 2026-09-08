import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "HUONG_DAN_SU_DUNG_SUPER_POST_DAY_DU.md"
OUTPUT = ROOT / "docs" / "HUONG_DAN_SU_DUNG_SUPER_POST_DAY_DU.docx"
LOGO = ROOT / "public" / "assets" / "branding" / "super-post-king-logo.png"

# compact_reference_guide with named Super Post brand overrides.
FONT = "Calibri"
INK = "1F2937"
MUTED = "6B7280"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "F15A24"
BROWN = "29190F"
LIGHT_ORANGE = "FFF3E8"
LIGHT_BLUE = "EFF6FF"
BORDER = "E5E7EB"


def set_run_font(run, size=None, color=None, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)


def add_left_border(paragraph, color=ORANGE, size=20):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(document, marker="%1.", start=1):
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(item.get(qn("w:numId")))
        for item in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)

    for level in range(2):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start_node = OxmlElement("w:start")
        start_node.set(qn("w:val"), str(start))
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if marker == "%1." else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), marker if level == 0 else marker.replace("%1", "%2"))
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        marker_pos = 269 + level * 360
        text_pos = 540 + level * 360
        tab.set(qn("w:pos"), str(text_pos))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(text_pos))
        ind.set(qn("w:hanging"), str(text_pos - marker_pos))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start_node, num_fmt, lvl_text, suffix, p_pr])
        abstract.append(lvl)

    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id, level=0):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_inline_runs(paragraph, text):
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True, color=INK)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, color=INK)


def configure_style(style, size, color, before, after, line=1.0, bold=True):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line
    style.paragraph_format.keep_with_next = True


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    configure_style(document.styles["Heading 1"], 16, ORANGE, 18, 10, 1.0)
    configure_style(document.styles["Heading 2"], 13, BROWN, 14, 7, 1.0)
    configure_style(document.styles["Heading 3"], 12, DARK_BLUE, 10, 5, 1.0)
    document.styles["Heading 1"].paragraph_format.page_break_before = True


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_running_header_footer(section):
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    left = paragraph.add_run("SUPER POST KING")
    set_run_font(left, size=8.5, bold=True, color=ORANGE)
    right = paragraph.add_run("    HƯỚNG DẪN SỬ DỤNG")
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.paragraph_format.space_before = Pt(0)
    label = footer_paragraph.add_run("Trang ")
    set_run_font(label, size=9, color=MUTED)
    add_page_field(footer_paragraph)


def add_cover(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(20)
    if LOGO.exists():
        run = paragraph.add_run()
        shape = run.add_picture(str(LOGO), width=Inches(1.15))
        shape._inline.docPr.set("descr", "Logo Super Post King")
        shape._inline.docPr.set("title", "Super Post King")

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(12)
    run = kicker.add_run("KING AUTOMATION AI")
    set_run_font(run, size=10, bold=True, color=ORANGE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("HƯỚNG DẪN SỬ DỤNG\nSUPER POST KING")
    set_run_font(run, size=29, bold=True, color=BROWN)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(
        "Đăng bài đa nền tảng • Quản lý profile • AI nội dung • Đồng bộ dữ liệu"
    )
    set_run_font(run, size=13, color=DARK_BLUE)

    lead = document.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.space_after = Pt(18)
    run = lead.add_run(
        "Hướng dẫn từng bước, dễ đọc và có hình minh họa trực tiếp trên tool"
    )
    set_run_font(run, size=11, italic=True, color=MUTED)

    cover_image = ROOT / "docs" / "images" / "huong-dan-toan-bo" / "01-dashboard.png"
    if cover_image.exists():
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = picture.add_run()
        shape = run.add_picture(str(cover_image), width=Inches(6.25))
        shape._inline.docPr.set("descr", "Dashboard tổng quan Super Post King")
        shape._inline.docPr.set("title", "Dashboard Super Post King")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(16)
    meta.paragraph_format.space_after = Pt(0)
    run = meta.add_run("Phiên bản tài liệu: 17/08/2026")
    set_run_font(run, size=9.5, color=MUTED)


def extract_body_lines():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    for index, line in enumerate(lines):
        if line.strip() == "---":
            return lines[index + 1 :]
    return lines


def extract_chapters(lines):
    chapters = []
    for line in lines:
        if line.startswith("## "):
            title = line[3:].strip()
            if re.match(r"\d+\.", title):
                chapters.append(title)
    return chapters


def add_contents_page(document, chapters):
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = False
    heading.add_run("Mục lục nhanh")
    intro = document.add_paragraph(
        "Tìm đúng tên mục bên dưới để tra cứu nhanh theo menu đang sử dụng."
    )
    intro.paragraph_format.space_after = Pt(10)
    bullet_id = add_numbering_definition(document, "•")
    for chapter in chapters:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        apply_numbering(paragraph, bullet_id)
        run = paragraph.add_run(chapter)
        set_run_font(run, size=9.5, bold=True, color=BROWN)


def image_width(path):
    with Image.open(path) as image:
        width, height = image.size
    ratio = width / max(height, 1)
    if width >= 1200:
        return 6.5
    if ratio >= 2.4:
        return 6.25
    if ratio >= 1.45:
        return 5.4
    return 4.65


def add_image(document, caption, relative_path):
    path = ROOT / "docs" / relative_path
    if not path.exists():
        warning = document.add_paragraph()
        warning.add_run(f"[Thiếu ảnh: {relative_path}]")
        shade_paragraph(warning, LIGHT_ORANGE)
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(image_width(path)))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)

    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(9)
    cap.paragraph_format.keep_with_next = False
    run = cap.add_run(f"Hình: {caption}")
    set_run_font(run, size=9, italic=True, color=MUTED)


def add_note(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.05)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    shade_paragraph(paragraph, LIGHT_ORANGE)
    add_left_border(paragraph)
    add_inline_runs(paragraph, text)


def flush_paragraph(document, buffer):
    if not buffer:
        return
    text = " ".join(part.strip() for part in buffer).strip()
    if text:
        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, text)
    buffer.clear()


def add_markdown_body(document, lines):
    paragraph_buffer = []
    current_list_type = None
    current_num_id = None

    def reset_list():
        nonlocal current_list_type, current_num_id
        current_list_type = None
        current_num_id = None

    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()

        if not stripped:
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            index += 1
            continue

        image_match = re.match(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image_match:
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            add_image(document, image_match.group(1), image_match.group(2))
            index += 1
            continue

        if stripped == "---":
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            index += 1
            continue

        if stripped.startswith("## "):
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            document.add_paragraph(stripped[3:].strip(), style="Heading 1")
            index += 1
            continue

        if stripped.startswith("### "):
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            document.add_paragraph(stripped[4:].strip(), style="Heading 2")
            index += 1
            continue

        if stripped.startswith("#### "):
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            document.add_paragraph(stripped[5:].strip(), style="Heading 3")
            index += 1
            continue

        if stripped.startswith("> "):
            flush_paragraph(document, paragraph_buffer)
            reset_list()
            note_parts = [stripped[2:].strip()]
            index += 1
            while index < len(lines) and lines[index].strip().startswith(">"):
                note_parts.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_note(document, " ".join(note_parts))
            continue

        number_match = re.match(r"^(\s*)(\d+)\.\s+(.+)$", raw)
        if number_match:
            flush_paragraph(document, paragraph_buffer)
            level = 1 if len(number_match.group(1)) >= 3 else 0
            if current_list_type != "number":
                current_num_id = add_numbering_definition(document, "%1.")
                current_list_type = "number"
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, current_num_id, level)
            add_inline_runs(paragraph, number_match.group(3))
            index += 1
            continue

        bullet_match = re.match(r"^(\s*)-\s+(.+)$", raw)
        if bullet_match:
            flush_paragraph(document, paragraph_buffer)
            level = 1 if len(bullet_match.group(1)) >= 3 else 0
            text = bullet_match.group(2)
            is_check = text.startswith("[ ] ")
            if is_check:
                text = text[4:]
            list_type = "check" if is_check else "bullet"
            if current_list_type != list_type:
                marker = "☐" if is_check else "•"
                current_num_id = add_numbering_definition(document, marker)
                current_list_type = list_type
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, current_num_id, level)
            add_inline_runs(paragraph, text)
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph(document, paragraph_buffer)


def set_modern_word_compatibility(document):
    settings = document.settings._element
    for item in settings.findall(f".//{qn('w:compatSetting')}"):
        if item.get(qn("w:name")) == "compatibilityMode":
            item.set(qn("w:val"), "15")
            return
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    item = OxmlElement("w:compatSetting")
    item.set(qn("w:name"), "compatibilityMode")
    item.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    item.set(qn("w:val"), "15")
    compat.append(item)


def build_document():
    lines = extract_body_lines()
    chapters = extract_chapters(lines)
    document = Document()
    configure_document(document)
    add_cover(document)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    add_running_header_footer(body_section)
    add_contents_page(document, chapters)
    add_markdown_body(document, lines)

    set_modern_word_compatibility(document)
    document.core_properties.title = "Hướng dẫn sử dụng Super Post King"
    document.core_properties.subject = "Hướng dẫn người dùng đầy đủ có hình minh họa"
    document.core_properties.author = "King Automation AI"
    document.core_properties.last_modified_by = "King Automation AI"
    document.core_properties.keywords = (
        "Super Post King, hướng dẫn sử dụng, đăng bài đa nền tảng, AI, Google Sheets"
    )
    document.core_properties.comments = ""
    document.core_properties.category = "User Guide"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
